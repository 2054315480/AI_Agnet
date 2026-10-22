#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
生成多格式测试语料 + 知识图谱素材（可重复执行，随机种子固定）。

产出目录：data/data/knowledge_base/{保险,物流,商城,政企服务,测试边界,图谱素材}
覆盖格式：md / docx / xlsx / csv / json / txt
用途：
  1) 把检索/评测/图谱测试从"3 域 7 文档"扩到接近真实业务规模；
  2) xlsx/csv/json/txt 当前不被解析器支持 —— 这批文件就是"待支持格式"的验收样本。

依赖：pip install python-docx openpyxl
"""
import io, os, json, csv, random

ROOT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                    "data", "data", "knowledge_base")
random.seed(20261022)
MADE = []


def full(rel):
    p = os.path.join(ROOT, rel)
    os.makedirs(os.path.dirname(p), exist_ok=True)
    return p


def W(rel, text):
    p = full(rel)
    io.open(p, "w", encoding="utf-8", newline="\n").write(text)
    MADE.append(p)
    return p


def SAVE(wb, rel):
    p = full(rel)
    wb.save(p)
    MADE.append(p)
    return p


def gen_insurance():
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from docx import Document

    items = ["意外险", "百万医疗险", "重疾险", "车险", "旅行险", "家财险"]
    asks = ["理赔要准备哪些材料", "多久能赔下来", "能不能线上申请", "需要原件吗", "异地理赔怎么办", "免赔额是多少"]
    md = "# 保险产品与理赔说明\n\n> 多格式语料之一，验证 Markdown 解析与切片。\n\n"
    md += "\n\n".join(
        "Q: %s的%s？\nA: %s类保单的%s，需在出险后 %d 个工作日内提交申请；材料齐全后 %d 个工作日内完成审核。"
        % (it, a, it, a, random.choice([3, 5, 7, 10]), random.choice([3, 5, 7]))
        for it in items for a in asks)
    md += "\n\n## 理赔时效总表\n\n| 险种 | 报案时限 | 审核时限 | 赔付时限 |\n| --- | --- | --- | --- |\n"
    md += "".join("| %s | %d个工作日 | %d个工作日 | %d个工作日 |\n"
                  % (it, random.choice([3, 5]), random.choice([3, 5]), random.choice([5, 7, 10])) for it in items)
    W("保险/保险产品与理赔.md", md)

    wb = Workbook(); ws = wb.active; ws.title = "理赔材料清单"
    ws["A1"] = "理赔材料清单（按险种）"; ws["A1"].font = Font(bold=True, size=13); ws.merge_cells("A1:D1")
    ws.append([]); ws.append(["险种", "必需材料", "选填材料", "备注"])
    for col, val in (("A", "险种"), ("B", "必需材料"), ("C", "选填材料"), ("D", "备注")):
        ws["%s4" % col] = val; ws.merge_cells("%s4:%s5" % (col, col))
    mats = {"意外险": ["身份证明", "出险证明", "医疗发票"], "车险": ["行驶证", "驾驶证", "事故认定书"],
            "重疾险": ["诊断证明", "病理报告", "住院病历"], "旅行险": ["行程单", "延误证明", "护照复印件"]}
    r = 6
    for k, v in mats.items():
        ws.cell(r, 1, k); ws.cell(r, 2, "、".join(v)); ws.cell(r, 3, "银行卡复印件"); ws.cell(r, 4, "原件核对后可退回"); r += 1
    ws2 = wb.create_sheet("理赔时效"); ws2.append(["险种", "报案时限(工作日)", "审核时限(工作日)", "赔付时限(工作日)", "备注"])
    for it in items:
        ws2.append([it, random.choice([3, 5]), random.choice([3, 5]), random.choice([5, 7, 10]), "以保单载明为准"])
    ws3 = wb.create_sheet("边缘案例"); ws3.append(["场景", "数值", "说明", "公式列"])
    for i in range(1, 13):
        ws3.append(["空值测试%d" % i, None if i % 3 == 0 else i * 10, "含空单元格与长文本" + "长" * (i * 3), "=B%d*2" % (i + 1)])
    ws3["F1"] = "合并表头陷阱"; ws3.merge_cells("F1:H1"); ws3["F2"] = "跨列表头"; ws3.merge_cells("F2:H2")
    SAVE(wb, "保险/理赔材料清单.xlsx")

    p = full("保险/保费与责任免除.csv")
    with io.open(p, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f); w.writerow(["险种", "年龄段", "年保费(元)", "保额(万元)", "责任免除条目", "是否可线上投保"])
        for it in items:
            for age in ["18-30", "31-45", "46-60"]:
                w.writerow([it, age, random.randint(120, 3200), random.choice([50, 100, 200, 300]),
                            random.choice(["既往症", "高危运动", "酒驾", "战争", "等待期内出险"]), random.choice(["是", "否"])])
    MADE.append(p)

    doc = Document(); doc.add_heading("保险条款摘要（DOCX 语料）", 0)
    doc.add_paragraph("用于验证 DOCX 解析：标题层级、列表、表格与长段落。")
    for it in items[:4]:
        doc.add_heading(it, level=1)
        doc.add_paragraph("意外身故/伤残、意外医疗、住院津贴（以条款为准）", style="List Bullet")
        doc.add_paragraph("等待期与免赔额详见保单。" + "保险责任范围以保险合同载明为准，" * 12)
        t = doc.add_table(rows=1, cols=3); t.style = "Table Grid"
        h = t.rows[0].cells; h[0].text = "项目"; h[1].text = "约定"; h[2].text = "备注"
        for k, v in [("保险期间", "1年"), ("缴费方式", "年缴"), ("犹豫期", "15天")]:
            c = t.add_row().cells; c[0].text = k; c[1].text = v; c[2].text = "以保单为准"
    p = full("保险/保险条款摘要.docx"); doc.save(p); MADE.append(p)


def gen_logistics():
    from openpyxl import Workbook
    md = "# 配送时效与运费说明\n\n"
    for c in ["北京", "上海", "广州", "成都", "西安", "乌鲁木齐", "拉萨"]:
        md += "## %s\n\n- 标准快递：%d 天送达，首重 %d 元\n- 次日达：%d 天，加收 %d 元\n- 偏远地区：%d 天，暂不支持次日达\n\n" % (
            c, random.choice([2, 3, 4]), random.randint(8, 15), random.choice([1, 2]), random.randint(5, 20), random.choice([5, 7, 10]))
    md += "## 常见问题\n\n" + "".join(
        "Q: 运单长时间不更新怎么办（%02d）？\nA: 超过 48 小时未更新可申请查件，客服 24 小时内反馈；确认丢件后按保价金额赔付。\n\n" % i
        for i in range(1, 26))
    W("物流/配送时效与运费.md", md)

    cases = [{"caseId": "EX%03d" % i, "type": random.choice(["延误", "丢件", "破损", "拒收", "地址错误"]),
              "sla": "%d小时" % random.choice([24, 48, 72]),
              "steps": ["核实运单", "联系网点", "回访客户", "补偿或重发"],
              "compensation": {"mode": random.choice(["按保价", "按运费倍数", "协商"]), "cap": random.randint(100, 2000)}}
             for i in range(1, 13)]
    W("物流/运单异常处理.json", json.dumps({"version": "2026-10", "cases": cases}, ensure_ascii=False, indent=2))

    wb = Workbook(); ws = wb.active; ws.title = "网点覆盖"
    ws.append(["省份", "城市", "网点名称", "网点编码", "是否支持次日达", "覆盖区县数"])
    provs = ["广东", "江苏", "浙江", "山东", "河南", "四川", "湖北", "福建", "陕西", "云南"]
    for i in range(1, 81):
        pv = provs[i % len(provs)]
        ws.append([pv, pv + "市%d区" % (i % 20 + 1), "%s转运中心%02d" % (pv, i), "ST%05d" % (10000 + i),
                   random.choice(["是", "否"]), random.randint(3, 30)])
    SAVE(wb, "物流/网点覆盖表.xlsx")
    W("物流/末端服务说明.txt", "末端服务与代收说明（TXT 纯文本语料）\n" + "=" * 40 + "\n" +
      "\n".join("第%d条 快递员应在投递前电话联系收件人；未经同意不得投放快递柜。" % i for i in range(1, 41)))


def gen_mall():
    from openpyxl import Workbook
    cats = {"数码": ["蓝牙耳机", "智能手表", "移动电源"], "家电": ["空气炸锅", "扫地机器人", "净水器"],
            "服饰": ["运动鞋", "冲锋衣", "速干T恤"], "食品": ["坚果礼盒", "咖啡豆", "蜂蜜"]}
    wb = Workbook(); ws = wb.active; ws.title = "商品目录"
    ws.append(["商品编码", "商品名称", "类目", "型号", "关键参数", "价格(元)", "库存", "保修期(月)"])
    n = 1
    for cat, names in cats.items():
        for nm in names:
            for spec in range(1, 6):
                ws.append(["SP%04d" % n, "%s %s款%d" % (cat, nm, spec), cat, "%s-%d" % (cat[:2].upper(), spec),
                           "参数：" + random.choice(["降噪", "防水", "快充", "静音", "大容量"]),
                           random.randint(39, 4999), random.randint(0, 999), random.choice([0, 6, 12, 24])]); n += 1
    ws2 = wb.create_sheet("价格与促销"); ws2.append(["商品编码", "原价", "促销价", "促销类型", "生效日期", "失效日期"])
    for i in range(1, 41):
        ws2.append(["SP%04d" % i, random.randint(100, 5000), random.randint(50, 4000),
                    random.choice(["满减", "折扣", "秒杀", "买赠"]), "2026-10-01", "2026-11-11"])
    ws3 = wb.create_sheet("FAQ"); ws3.append(["问题", "答案", "适用类目"])
    for i in range(1, 31):
        ws3.append(["商品能否指定配送时间（%d）？" % i, "支持在结算页选择时段，以实际可选项为准。", random.choice(list(cats.keys()))])
    SAVE(wb, "商城/商品目录.xlsx")

    md = "# 促销规则说明\n\n" + "".join(
        "## 规则%02d：%s\n\n- 适用类目：%s\n- 门槛：满 %d 元\n- 优惠：减 %d 元\n- 叠加限制：%s\n\n" % (
            i, random.choice(["满减", "折扣", "秒杀", "买赠"]), random.choice(list(cats.keys())),
            random.choice([99, 199, 299, 599]), random.choice([10, 30, 50, 100]),
            random.choice(["不可叠加", "可叠加一张券", "仅限单品"])) for i in range(1, 21))
    W("商城/促销规则.md", md)

    states = {"订单状态": ["待支付", "已支付", "已发货", "运输中", "已签收", "已完成", "已取消", "退款中", "已退款"],
              "流转": [{"from": a, "to": b, "trigger": t} for a, b, t in
                     [("待支付", "已支付", "支付成功"), ("已支付", "已发货", "商家发货"), ("已发货", "运输中", "揽收"),
                      ("运输中", "已签收", "签收"), ("已签收", "已完成", "超7天自动完成"), ("待支付", "已取消", "超时未支付"),
                      ("已完成", "退款中", "申请售后"), ("退款中", "已退款", "退款到账")]]}
    W("商城/订单状态说明.json", json.dumps(states, ensure_ascii=False, indent=2))


def gen_gov():
    faq = ["# 常见问题 200 条（批量语料）\n\n"] + [
        "Q%03d: %s？\nA: %s，办理时限 %d 个工作日，可通过线上或线下渠道提交，所需材料见对应业务指南。\n\n" % (
            i, random.choice(["如何办理", "需要什么材料", "多久生效", "能否代办", "收费多少", "如何查询进度"]),
            random.choice(["携带本人身份证", "提供单位证明", "线上实名认证后即可", "需先完成前置业务"]),
            random.choice([1, 3, 5, 7, 15])) for i in range(1, 201)]
    W("政企服务/常见问题200条.md", "".join(faq))

    p = full("政企服务/版本变更记录.csv")
    with io.open(p, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f); w.writerow(["变更版本", "生效日期", "变更类型", "变更条目", "影响范围", "是否向下兼容"])
        for v, d in [("v1.0", "2025-06-01"), ("v1.5", "2026-01-01"), ("v2.0", "2026-09-01")]:
            for i in range(1, 13):
                w.writerow([v, d, random.choice(["新增", "修订", "废止"]),
                            "第%d条 %s" % (i, random.choice(["时效调整", "材料简化", "费率变更", "渠道扩展"])),
                            random.choice(["全部", "个人业务", "企业业务"]), random.choice(["是", "否"])])
    MADE.append(p)
    W("政企服务/服务承诺与时效.md",
      "# 服务承诺\n\n- 首次响应：15 分钟内\n- 办结时限：5 个工作日\n- 超时补偿：每日按 20 元补偿\n\n"
      "> 历史版本 v1.0 的承诺为 30 分钟响应，v2.0 起调整。\n")


def gen_edge():
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "合并单元格"
    for rr in range(1, 4):
        ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=5)
        ws.cell(rr, 1, "第%d层表头（跨5列合并）" % rr)
    ws.append(["子表头A", "子表头B", None, None, "子表头E"]); ws.append(["值1", "", "值3", None, "值5"])
    ws.merge_cells("A6:B7"); ws["A6"] = "纵向合并单元格"
    ws["C6"] = "换行\n测试"; ws["D6"] = " 前后空格 "; ws["E6"] = "=SUM(1,2)"
    for i in range(8, 28):
        ws.append([None if i % 4 == 0 else "行%d" % i, "空行与断档测试", i, None, i * 1.5])
    SAVE(wb, "测试边界/合并单元格陷阱.xlsx")

    wb = Workbook(); ws = wb.active; ws.title = "公式与空值"; ws.append(["项目", "数量", "单价", "金额(公式)", "备注"])
    for i in range(2, 32):
        ws.cell(i, 1, "项目%d" % i); ws.cell(i, 2, None if i % 5 == 0 else random.randint(1, 20))
        ws.cell(i, 3, random.randint(10, 500)); ws.cell(i, 4, "=B%d*C%d" % (i, i))
        ws.cell(i, 5, random.choice(["", "含,逗号", '含"引号"', "含\n换行", "极长备注" + "长" * 40]))
    SAVE(wb, "测试边界/公式与空值.xlsx")

    long_para = "长文本压力测试段落：" + "".join("本段用于验证固定窗口切片是否会把语义切碎，" for _ in range(30))
    W("测试边界/超长段落与噪声.md",
      "# 超长段落与噪声\n\n" + long_para + "\n\n## 含 HTML 与代码\n\n<div>HTML 块</div>\n\n```sql\nSELECT * FROM cs_vector_store;\n```\n\n"
      "## 特殊符号\n\nemoji 😀🚚✅、全角（）【】、数学 ∑≈、货币 ¥1,234.56、路径 /opt/ai-agent/data\n\n"
      "## 表格\n\n| 列1 | 列2 |\n| --- | --- |\n| 值 | 值 |\n")

    W("测试边界/重复与冲突问句.md",
      "# 重复与冲突问句\n\nQ: 退款多久到账？\nA: 3-5 个工作日。\n\nQ: 退款多久到账？\nA: 7-15 个工作日（旧版本，2025-01 前）。\n\n"
      "Q: 退换货政策是什么？\nA: 7 天无理由。\n\nQ: 退换货政策是什么？\nA: 15 天无理由（v2 政策，2026-09 起）。\n\n" +
      "\n".join("Q: 重复问题样本 %02d：如何查询订单？\nA: 在订单列表选择对应订单查看详情。" % i for i in range(1, 16)))

    rows = [["字段A", "字段B", "字段C"]] + [["含,逗号的%s" % i, '含"双引号"的值',
                                          "含\n换行的值" if i % 5 == 0 else "普通值"] for i in range(1, 31)]
    p = full("测试边界/特殊字符与乱码.csv")
    with io.open(p, "w", encoding="utf-8-sig", newline="") as f:
        csv.writer(f).writerows(rows)
    MADE.append(p)


def gen_kg_material():
    products = ["蓝牙耳机", "智能手表", "扫地机器人", "空气炸锅", "流量套餐", "宽带套餐"]
    issues = ["左耳无声", "续航缩短", "无法开机", "吸力下降", "频繁断电", "网速慢"]
    policies = ["七天无理由退货", "一年质保", "上门维修", "只换不修", "免运费退换"]
    channels = ["官网", "APP", "客服热线", "在线客服", "线下营业厅", "微信小程序"]
    acts = ["申请退款", "申请换货", "报修", "查询进度", "升级套餐", "转人工"]
    md = "# 图谱实体关系语料（供知识图谱抽取）\n\n" + "".join(
        "场景%02d：用户通过%s反馈「%s出现%s」，依据%s，客服引导其%s，承诺%d个工作日处理完成，超时可申请补偿。\n" % (
            i, channels[i % len(channels)], products[i % len(products)], issues[i % len(issues)],
            policies[i % len(policies)], acts[i % len(acts)], random.choice([1, 3, 5, 7])) for i in range(1, 61))
    W("图谱素材/实体关系语料.md", md)
    W("图谱素材/kg_dictionary_expansion.json", json.dumps({
        "说明": "KgExtractor 现有词典偏薄（单字/通用词污染图谱），这里给出扩展建议，可直接补进静态词典。",
        "entityTypes": {
            "PRODUCT": ["蓝牙耳机", "智能手表", "扫地机器人", "空气炸锅", "净水器", "移动电源", "流量套餐", "宽带套餐", "信用卡", "储蓄卡"],
            "PART": ["左耳", "右耳", "充电仓", "耳塞", "电池", "滤芯", "电机", "屏幕", "主板", "SIM卡"],
            "ISSUE": ["无声", "杂音", "续航缩短", "无法开机", "无法充电", "吸力下降", "漏水", "破损", "延迟发货", "错发", "少件"],
            "ACTION": ["申请退款", "申请换货", "申请退货", "报修", "上门维修", "查询进度", "修改地址", "取消订单", "升级套餐", "转人工"],
            "POLICY": ["七天无理由退货", "十五天换货", "一年质保", "两年质保", "只换不修", "免运费退换", "价保", "延保"],
            "CHANNEL": ["官网", "APP", "微信小程序", "在线客服", "客服热线", "线下营业厅", "快递员"],
            "DURATION": ["1个工作日", "3个工作日", "7个工作日", "15天", "24小时", "48小时"],
            "AMOUNT": ["免赔额", "保额", "运费", "保价金额", "退款金额"],
            "ORG": ["售后服务中心", "理赔部", "网点", "转运中心", "营业厅"]},
        "relationTypes": [
            {"name": "HAS_PART", "from": "PRODUCT", "to": "PART", "样例": "蓝牙耳机 -> 左耳"},
            {"name": "HAS_ISSUE", "from": "PRODUCT", "to": "ISSUE", "样例": "蓝牙耳机 -> 无声"},
            {"name": "COVERED_BY", "from": "ISSUE", "to": "POLICY", "样例": "无声 -> 一年质保"},
            {"name": "RESOLVED_BY", "from": "ISSUE", "to": "ACTION", "样例": "无声 -> 报修"},
            {"name": "SUBMITTED_VIA", "from": "ACTION", "to": "CHANNEL", "样例": "申请退款 -> 官网"},
            {"name": "WITHIN", "from": "ACTION", "to": "DURATION", "样例": "报修 -> 3个工作日"},
            {"name": "OWNED_BY", "from": "POLICY", "to": "ORG", "样例": "一年质保 -> 售后服务中心"}],
        "抽取流水线建议": ["规则词典（快、零成本，当前实现）",
                     "LLM 抽取严格 JSON + Schema 校验 + 失败重试 1 次",
                     "并行构建（8 线程）+ 死信表 + 降级为仅规则抽取"],
        "负面词表（应过滤的过宽实体）": ["修", "查", "换", "退", "买", "客服", "反馈", "手机", "电话", "问题", "服务", "处理"]
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    gen_insurance(); gen_logistics(); gen_mall(); gen_gov(); gen_edge(); gen_kg_material()
    print("生成 %d 个文件，总 %.1f KB" % (len(MADE), sum(os.path.getsize(m) for m in MADE) / 1024))
    from collections import Counter
    print("格式分布:", dict(Counter(os.path.splitext(m)[1].lower() for m in MADE)))
    print("根目录:", ROOT)
